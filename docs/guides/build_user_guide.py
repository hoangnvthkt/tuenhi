from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("huong-dan-su-dung-co-ban-tuenhi.docx")

# compact_reference_guide preset + editorial_cover first-page pattern.
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "0F766E"
MUTED = "64748B"
PALE_BLUE = "E8EEF5"
PALE_TEAL = "ECFDF5"
PALE_GOLD = "FFFBEB"
INK = "0F172A"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_font(run, size=None, color=INK, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT):
    table.autofit = False
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_borders(table, color="CBD5E1"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    ppr.append(keep)


def set_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, 9, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_text(doc, text, bold_prefix=None, after=6, color=INK, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, 11, color, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, 11, color, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, 11, color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_keep_with_next(p)
    return p


def new_decimal_numbering(doc):
    """Create a fresh real Word numbering instance so each procedure starts at 1."""
    numbering = doc.part.numbering_part.element
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    next_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "7")  # Built-in List Number abstract definition.
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return next_id


def set_numbering(paragraph, number_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(number_id))
    num_pr.extend([ilvl, num_id])
    ppr.append(num_pr)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_font(p.add_run(f"{index}. {item}"), 11)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_font(p.add_run(f"• {item}"), 11)


def add_note(doc, title, text, fill=PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    set_table_borders(table, "FCD34D" if fill == PALE_GOLD else "99F6E4")
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, 10.5, "92400E" if fill == PALE_GOLD else "115E59", bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text), 10.5, "713F12" if fill == PALE_GOLD else "134E4A")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, [2700, 6660])
    set_table_borders(table)
    for label, value in rows:
        cells = table.add_row().cells
        shade(cells[0], PALE_BLUE)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), 10.5, NAVY, bold=True)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_two_column_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [3600, 5760])
    set_table_borders(table)
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(header), 10.5, NAVY, bold=True)
    for left, right in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, (left, right)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_font(p.add_run(value), 10.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def new_page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("TUỆ NHI  |  HƯỚNG DẪN SỬ DỤNG CƠ BẢN")
    set_font(run, 8.5, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    set_page_field(footer)


def cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(100)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("SỔ TAY THAO TÁC"), 11, TEAL, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run("Hướng dẫn sử dụng cơ bản"), 29, NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    set_font(p.add_run("Ứng dụng Tuệ Nhi - Bán hàng & Kho"), 15, DARK_BLUE)

    add_note(
        doc,
        "Dành cho người mới",
        "Tài liệu tập trung vào các thao tác hằng ngày: quản lý hàng hóa, nhà cung cấp, nhập hàng và bán hàng.",
        PALE_TEAL,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    set_font(p.add_run("Phiên bản hướng dẫn nội bộ | 09/2026"), 10, MUTED, italic=True)
    new_page(doc)


def guide(doc):
    add_heading(doc, "1. Bắt đầu sử dụng", 1)
    add_text(doc, "Tuệ Nhi là ứng dụng nội bộ hỗ trợ một cửa hàng quản lý danh mục hàng hóa, nhập kho và bán hàng. Những nút/chức năng bạn nhìn thấy phụ thuộc vào quyền của tài khoản.")
    add_key_value_table(doc, [
        ("Hàng hóa", "Tra cứu tồn, giá bán và thông tin sản phẩm; người có quyền có thể thêm/sửa sản phẩm và nhóm hàng."),
        ("Nhiều hơn", "Nơi mở Nhà cung cấp, Nhập hàng, Kiểm kho, Trả hàng, Báo cáo và cấu hình (tùy quyền)."),
        ("Bán hàng", "Tìm hàng, lập giỏ, lưu tạm và thanh toán để tạo hóa đơn."),
        ("Hóa đơn", "Tra cứu giao dịch đã hoàn tất, in/tải PDF, đối soát và xử lý trả/hủy khi được cấp quyền."),
    ])
    add_note(doc, "Nguyên tắc an toàn", "Khi thao tác với tiền hoặc tồn kho, hãy kiểm tra lại hàng hóa, số lượng, giá và chứng từ trước khi bấm Ghi sổ hoặc Xác nhận thanh toán.")
    add_heading(doc, "Trước khi bắt đầu", 2)
    add_bullets(doc, [
        "Đảm bảo thiết bị có kết nối mạng. Khi ngoại tuyến, các lệnh lưu, gửi, ghi sổ và thanh toán có thể bị khóa.",
        "Dùng đúng tài khoản của mình. Nếu không thấy chức năng cần dùng, hãy liên hệ chủ cửa hàng để kiểm tra quyền.",
        "Không mở cùng một giỏ bán hàng ở nhiều tab. Nếu thấy thông báo giỏ đang được chỉnh sửa ở tab khác, chỉ tiếp tục tại một tab.",
    ])

    add_heading(doc, "2. Quản lý hàng hóa", 1)
    add_text(doc, "Mở Hàng hóa từ thanh điều hướng để tra cứu danh mục, giá bán hiện hành và số lượng tồn.")
    add_heading(doc, "2.1 Tìm và kiểm tra sản phẩm", 2)
    add_numbered(doc, [
        "Nhập tên, SKU hoặc thông tin cần tìm vào vùng lọc; có thể lọc theo nhóm hàng hoặc trạng thái tồn.",
        "Danh sách hiển thị ảnh đại diện của hàng hóa đã đặt ảnh chính để nhận diện nhanh hơn.",
        "Từ trang chi tiết, dùng Bán hàng hoặc Nhập hàng để đi thẳng sang nghiệp vụ tương ứng khi có quyền.",
    ])
    add_heading(doc, "2.2 Thêm hoặc sửa sản phẩm", 2)
    add_text(doc, "Người có quyền quản lý chọn Thêm sản phẩm hoặc Sửa sản phẩm. Điền các trường theo bảng dưới đây rồi bấm Lưu sản phẩm.")
    add_two_column_table(doc, ("Trường", "Cách dùng"), [
        ("SKU", "Mã nội bộ để nhận diện sản phẩm. Nên dùng một quy ước thống nhất và không trùng giữa các mặt hàng."),
        ("Mã vạch", "Nhập nếu hàng có mã vạch. Có thể dùng máy quét như bàn phím khi tìm/bán hàng."),
        ("Tên sản phẩm", "Tên rõ ràng, dễ tìm; nên kèm quy cách nếu cần phân biệt."),
        ("Nhóm hàng", "Phân loại để dễ tìm và theo dõi. Có thể tạo/quản lý nhóm tại Quản lý nhóm hàng."),
        ("Đơn vị tính", "Ví dụ: cái, hộp, chai, kg. Đây là đơn vị hiển thị khi nhập và bán."),
        ("Ngưỡng tồn tối thiểu", "Mức tồn cần theo dõi để chủ động nhập thêm."),
        ("Giá bán hiện hành", "Chỉ chủ cửa hàng mới có thể thay đổi giá bán."),
        ("Đang hoạt động", "Chỉ hàng đang hoạt động mới nên dùng cho nghiệp vụ mới. Không cần xóa hàng đã có lịch sử."),
    ])
    add_note(doc, "Mẹo", "Hãy tạo nhóm hàng và chuẩn hóa SKU trước khi thêm nhiều sản phẩm. Điều này giúp tìm hàng nhanh hơn ở màn Bán hàng và Nhập hàng.", PALE_TEAL)

    add_heading(doc, "3. Quản lý nhà cung cấp (NCC)", 1)
    add_text(doc, "Vào Nhiều hơn > Nhà cung cấp. Màn hình này dùng để lưu thông tin liên hệ và liên kết NCC với các lần nhập hàng đã ghi sổ.")
    add_heading(doc, "3.1 Thêm hoặc cập nhật NCC", 2)
    add_numbered(doc, [
        "Chọn Thêm nhà cung cấp; để sửa, tìm NCC rồi chọn Sửa.",
        "Điền mã NCC (nếu có), tên NCC, số điện thoại, email, địa chỉ và ghi chú.",
        "Giữ trạng thái Nhà cung cấp đang hoạt động nếu còn giao dịch với NCC này.",
        "Bấm Lưu nhà cung cấp và kiểm tra thông báo thành công.",
    ])
    add_heading(doc, "3.2 Tra cứu NCC", 2)
    add_bullets(doc, [
        "Tìm theo tên, mã hoặc số điện thoại.",
        "Mở trang chi tiết để xem thông tin liên hệ; nếu có quyền xem nhập hàng, bạn còn thấy mặt hàng cung cấp, số phiếu đã ghi sổ và lịch sử nhập.",
        "Từ trang NCC, chọn Lập phiếu nhập để tạo phiếu đã điền sẵn NCC; bạn vẫn cần kiểm tra và lưu phiếu.",
    ])
    add_note(doc, "Lưu ý", "Ngừng hoạt động NCC khi không còn giao dịch mới. Không nên xóa NCC đã phát sinh lịch sử vì sẽ khó tra cứu lại các phiếu nhập cũ.")

    add_heading(doc, "4. Nghiệp vụ mua hàng - nhập kho", 1)
    add_text(doc, "Vào Nhiều hơn > Nhập hàng. Mỗi lần nhận hàng, hãy lập một phiếu nhập riêng để tồn kho và giá nhập được ghi nhận đúng.")
    add_heading(doc, "4.1 Lập phiếu nhập", 2)
    add_numbered(doc, [
        "Chọn lập phiếu nhập mới.",
        "Chọn Nhà cung cấp (có thể để trống nếu chưa xác định), kiểm tra Ngày nhận và nhập Ghi chú khi cần.",
        "Gõ tên, SKU hoặc mã vạch vào ô Sản phẩm để tìm và chọn hàng. Mỗi sản phẩm chỉ được chọn một lần trong cùng phiếu.",
        "Nhập số lượng nhận thực tế và Đơn giá nhập cho từng dòng; cả hai luôn phải lớn hơn 0. Người lập phiếu có quyền nhập giá có thể tự nhập giá, không cần chuyển phiếu cho owner.",
        "Bấm Lưu nháp để lưu lại và tiếp tục xử lý sau.",
    ])
    add_heading(doc, "4.2 Trạng thái và thao tác chính", 2)
    add_two_column_table(doc, ("Thao tác", "Ý nghĩa / khi dùng"), [
        ("Lưu nháp", "Lưu phiếu đang soạn. Phiếu nháp chưa làm thay đổi tồn kho."),
        ("Nhập từ Excel", "Tải file mẫu, điền đúng 3 cột SKU, Số lượng nhận và Đơn giá nhập; chọn file để xem trước. Hệ thống chỉ nạp vào phiếu khi toàn bộ dòng đều hợp lệ."),
        ("Ghi sổ", "Xác nhận phiếu; số lượng được ghi nhận vào tồn kho và giá nhập của từng dòng được áp dụng."),
        ("Đảo phiếu", "Thao tác có kiểm soát để xử lý một phiếu đã ghi sổ khi có quyền. Không tự tạo phiếu thay thế nếu chưa hiểu rõ lý do."),
    ])
    add_note(doc, "Nhập Excel an toàn", "Không sửa tên sheet hoặc cột của file mẫu. Nếu bảng xem trước báo lỗi SKU, số lượng, giá hoặc hàng bị trùng, hãy sửa toàn bộ lỗi trong file rồi chọn lại file. Không có dòng nào được nạp một phần.", PALE_TEAL)
    add_note(doc, "Điểm kiểm tra trước khi Ghi sổ", "Đối chiếu phiếu giao hàng: đúng NCC, đúng ngày nhận, đúng mặt hàng, đúng số lượng và giá nhập. Ghi sổ là thao tác ảnh hưởng trực tiếp đến tồn kho.")
    add_heading(doc, "4.3 Tra cứu sau khi nhập", 2)
    add_bullets(doc, [
        "Mở danh sách Phiếu nhập hàng để tìm theo số phiếu/trạng thái và xem lại chứng từ.",
        "Từ chi tiết sản phẩm hoặc NCC, xem Lịch sử nhập để kiểm tra những phiếu đã ghi sổ liên quan.",
        "Nếu có sai khác thực tế sau nhận hàng, trao đổi với người có quyền trước khi ghi sổ/đảo phiếu; không tự điều chỉnh bằng một phiếu không có căn cứ.",
    ])

    add_heading(doc, "5. Nghiệp vụ bán hàng", 1)
    add_text(doc, "Mở Bán hàng để tạo giỏ, áp dụng thông tin khách/kênh bán và thanh toán. Chỉ hoàn tất thanh toán khi hàng và số tiền đã được đối chiếu.")
    add_heading(doc, "5.1 Lập giỏ hàng", 2)
    add_numbered(doc, [
        "Tìm sản phẩm theo tên, SKU hoặc quét mã vạch bằng máy quét kết nối như bàn phím.",
        "Chọn hàng để thêm vào giỏ; chỉnh số lượng hoặc bỏ dòng không cần thiết.",
        "Chọn kênh bán và khách hàng nếu nghiệp vụ yêu cầu; với khách lẻ có thể để mặc định.",
        "Nhập giảm giá hoặc ghi chú nếu được cấp quyền, sau đó kiểm tra tiền hàng, giảm giá và tổng thanh toán.",
        "Nếu chưa thanh toán, bấm Lưu tạm để lưu nháp. Có thể quay lại hoàn thiện sau.",
    ])
    add_heading(doc, "5.2 Thanh toán và tạo hóa đơn", 2)
    add_numbered(doc, [
        "Bấm Thanh toán và chọn Tiền mặt hoặc Chuyển khoản.",
        "Với chuyển khoản, chuẩn bị ảnh chứng từ đúng định dạng JPEG, PNG hoặc WebP, dung lượng tối đa 5 MiB.",
        "Đối chiếu lại khách, mặt hàng, số lượng, giảm giá và tổng tiền.",
        "Bấm Xác nhận. Khi hoàn tất, hệ thống tạo hóa đơn và ghi nhận giao dịch theo phương thức đã chọn.",
    ])
    add_note(doc, "Không chắc kết quả?", "Nếu mạng chập chờn hoặc xuất hiện trạng thái chờ đối soát, không bấm thanh toán lặp lại. Vào hóa đơn để dùng Đối soát lại giao dịch hoặc báo người quản lý kèm mã yêu cầu nếu có.")
    add_heading(doc, "5.3 Sau khi bán", 2)
    add_bullets(doc, [
        "Vào Hóa đơn để tìm giao dịch đã hoàn tất theo thời gian/trạng thái.",
        "Mở hóa đơn để in nhiệt, tải PDF hoặc kiểm tra phương thức, trạng thái thanh toán và người thực hiện.",
        "Dùng Đơn mới để bắt đầu giao dịch tiếp theo sau khi đã hoàn tất hóa đơn hiện tại.",
    ])

    add_heading(doc, "6. Trả hàng, hủy hóa đơn và đối soát", 1)
    add_text(doc, "Các thao tác này chỉ hiện khi tài khoản có quyền và hóa đơn còn đủ điều kiện xử lý.")
    add_two_column_table(doc, ("Tình huống", "Cách xử lý an toàn"), [
        ("Khách trả một phần/toàn bộ", "Mở hóa đơn và chọn Tạo yêu cầu trả hàng. Chọn đúng dòng hàng, số lượng và lý do; thực hiện kiểm nhận/hoàn tiền theo quyền."),
        ("Cần hủy hóa đơn", "Chỉ hủy khi có lý do chính đáng. Nhập lý do hủy, kiểm tra lại đúng hóa đơn rồi bấm Hủy hóa đơn. Không hủy để sửa lỗi nhỏ nếu quy trình nội bộ yêu cầu trả hàng/điều chỉnh khác."),
        ("Giao dịch chờ đối soát", "Mở phần Đối soát thanh toán trên hóa đơn, dùng Đối soát lại giao dịch và chờ trạng thái cập nhật. Tránh tạo hóa đơn mới cho cùng một lần thanh toán."),
    ])

    add_heading(doc, "7. Checklist cuối ca", 1)
    add_bullets(doc, [
        "Kiểm tra không còn giỏ bán hàng hoặc phiếu nhập nháp chưa rõ người phụ trách.",
        "Đối chiếu các hóa đơn chuyển khoản với chứng từ ảnh và trạng thái đã ghi nhận.",
        "Kiểm tra các phiếu nhập trong ngày: NCC, hàng, số lượng, giá và trạng thái ghi sổ.",
        "Báo ngay cho quản lý nếu phát hiện tồn kho chênh lệch, thanh toán chờ đối soát hoặc thao tác ghi sổ nhầm.",
        "Đăng xuất khi dùng máy tính dùng chung.",
    ])

    new_page(doc)
    add_heading(doc, "8. Lỗi thường gặp", 1)
    add_two_column_table(doc, ("Dấu hiệu", "Nên làm gì"), [
        ("Không thấy nút Thêm/Sửa/Ghi sổ", "Khả năng cao tài khoản chưa có quyền. Kiểm tra đúng tài khoản và liên hệ chủ cửa hàng."),
        ("Không lưu được", "Kiểm tra kết nối mạng, các ô bắt buộc và thông báo lỗi bên dưới trường nhập. Đừng đóng trang trước khi ghi nhận nội dung lỗi."),
        ("Không thể thêm sản phẩm vào phiếu nhập", "Kiểm tra sản phẩm có đang hoạt động hay không và sản phẩm đó chưa xuất hiện trong cùng phiếu."),
        ("Không nạp được file Excel", "Dùng đúng file mẫu; kiểm tra đủ 3 cột SKU, Số lượng nhận, Đơn giá nhập và sửa mọi dòng có lỗi trong bảng xem trước."),
        ("Giỏ bán hàng đang bị khóa", "Giỏ đang được chỉnh sửa ở tab khác. Quay lại tab đang mở hoặc chọn Tiếp tục ở tab này nếu bạn chắc chắn sẽ dùng tab hiện tại."),
        ("Không chắc đã thanh toán thành công", "Không thanh toán lại ngay. Tìm hóa đơn, xem trạng thái thanh toán và dùng đối soát khi cần."),
    ])



def build():
    doc = Document()
    configure(doc)
    cover(doc)
    guide(doc)
    doc.core_properties.title = "Hướng dẫn sử dụng cơ bản Tuệ Nhi"
    doc.core_properties.subject = "Quản lý hàng hóa, nhà cung cấp và nghiệp vụ mua bán"
    doc.core_properties.author = "Tuệ Nhi"
    doc.save(OUT)


if __name__ == "__main__":
    build()
